from io import BytesIO
import json

from docx import Document

from docx.enum.section import (
    WD_ORIENT,
)

from docx.shared import (
    Cm,
    Pt,
)


class ReportDOCXService:

    @staticmethod
    def _text(
        value,
    ) -> str:

        if value is None:
            return "-"

        if isinstance(
            value,
            (dict, list),
        ):

            return json.dumps(
                value,
                ensure_ascii=False,
                default=str,
            )

        return str(value)


    @staticmethod
    def _origin(
        item: dict,
    ) -> str:

        metadata = (
            item.get(
                "metadata"
            )
            or {}
        )

        provenance = (
            item.get(
                "provenance"
            )
            or {}
        )


        if metadata.get(
            "data_origin"
        ):

            return str(
                metadata[
                    "data_origin"
                ]
            )


        if metadata.get(
            "ingestion_type"
        ):

            return str(
                metadata[
                    "ingestion_type"
                ]
            )


        if provenance.get(
            "source_name"
        ):

            return str(
                provenance[
                    "source_name"
                ]
            )


        if provenance.get(
            "source_type"
        ):

            return str(
                provenance[
                    "source_type"
                ]
            )


        return "-"


    @staticmethod
    def _configure_table(
        table,
        font_size: int = 8,
    ) -> None:

        table.style = (
            "Table Grid"
        )

        for row in (
            table.rows
        ):

            for cell in (
                row.cells
            ):

                for paragraph in (
                    cell.paragraphs
                ):

                    for run in (
                        paragraph.runs
                    ):

                        run.font.size = (
                            Pt(
                                font_size
                            )
                        )


    def generate(
        self,
        report: dict,
    ) -> bytes:

        document = (
            Document()
        )


        # =================================================
        # PAGE
        # =================================================

        section = (
            document.sections[
                0
            ]
        )

        section.orientation = (
            WD_ORIENT.LANDSCAPE
        )

        section.page_width = (
            Cm(29.7)
        )

        section.page_height = (
            Cm(21.0)
        )

        section.top_margin = (
            Cm(1.4)
        )

        section.bottom_margin = (
            Cm(1.4)
        )

        section.left_margin = (
            Cm(1.4)
        )

        section.right_margin = (
            Cm(1.4)
        )


        # =================================================
        # TITLE
        # =================================================

        document.add_heading(
            (
                "Reporte del Gemelo "
                "Digital Agrícola"
            ),
            level=0,
        )


        document.add_paragraph(
            (
                "Estado operativo, evidencia, "
                "recomendaciones, conocimiento "
                "del agricultor y gobernanza."
            )
        )


        parcel = (
            report.get(
                "parcel",
                {},
            )
        )


        # =================================================
        # 1. PARCEL
        # =================================================

        document.add_heading(
            "1. Parcela",
            level=1,
        )


        table = (
            document.add_table(
                rows=1,
                cols=2,
            )
        )


        table.rows[
            0
        ].cells[
            0
        ].text = "Campo"

        table.rows[
            0
        ].cells[
            1
        ].text = "Valor"


        parcel_values = {

            "ID":
                parcel.get(
                    "id"
                ),

            "Nombre":
                parcel.get(
                    "name"
                ),

            "Cultivo":
                parcel.get(
                    "crop_type"
                ),

            "Área (ha)":
                parcel.get(
                    "area_ha"
                ),

            "Fecha de generación":
                report.get(
                    "generated_at"
                ),
        }


        for key, value in (
            parcel_values.items()
        ):

            cells = (
                table.add_row()
                .cells
            )

            cells[
                0
            ].text = str(
                key
            )

            cells[
                1
            ].text = (
                self._text(
                    value
                )
            )


        self._configure_table(
            table,
            9,
        )


        # =================================================
        # 2. SUMMARY
        # =================================================

        document.add_heading(
            "2. Resumen",
            level=1,
        )


        summary_table = (
            document.add_table(
                rows=1,
                cols=2,
            )
        )


        summary_table.rows[
            0
        ].cells[
            0
        ].text = "Indicador"

        summary_table.rows[
            0
        ].cells[
            1
        ].text = "Valor"


        labels = {

            "evidence_count":
                "Evidencias totales",

            "active_evidence_count":
                "Evidencias activas",

            "stale_evidence_count":
                "Evidencias desactualizadas",

            "farmer_knowledge_count":
                "Aportes del agricultor",

            "decision_trace_count":
                "Trazas de decisión",

            "active_consent_count":
                "Consentimientos activos",

            "sensor_count":
                "Sensores registrados",

            "active_sensor_count":
                "Sensores activos",

            "constraint_count":
                "Restricciones registradas",

            "hard_constraint_count":
                "Restricciones HARD",

            "latest_conflict_count":
                "Conflictos última decisión",

            "latest_feasible_scenario_count":
                "Escenarios factibles",

            "latest_rejected_scenario_count":
                "Escenarios rechazados",

            "ml_forecast_available":
                "Pronóstico ML complementario",
        }


        for key, value in (
            report.get(
                "summary",
                {},
            ).items()
        ):

            cells = (
                summary_table
                .add_row()
                .cells
            )

            cells[
                0
            ].text = (
                labels.get(
                    key,
                    key,
                )
            )

            cells[
                1
            ].text = (
                self._text(
                    value
                )
            )


        self._configure_table(
            summary_table,
            9,
        )


        # =================================================
        # 3. LATEST DECISION
        # =================================================

        document.add_heading(
            "3. Última decisión",
            level=1,
        )


        latest = (
            report.get(
                "latest_decision"
            )
        )


        if latest:

            recommendation = (
                latest.get(
                    "provisional_recommendation"
                )
                or {}
            )

            farmer_response = (
                latest.get(
                    "farmer_response"
                )
                or {}
            )


            decision_values = {

                "Trace ID":
                    latest.get(
                        "trace_id"
                    ),

                "Fecha":
                    latest.get(
                        "created_at"
                    ),

                "Riego recomendado (mm)":
                    recommendation.get(
                        "irrigation_mm"
                    ),

                "Score":
                    recommendation.get(
                        "score"
                    ),

                "Estado":
                    recommendation.get(
                        "status"
                    ),

                "Fuente":
                    recommendation.get(
                        "source"
                    ),

                "Razones":
                    " | ".join(
                        recommendation.get(
                            "reasons",
                            [],
                        )
                    ),

                "Respuesta agricultor":
                    (
                        farmer_response.get(
                            "response_type"
                        )
                        if farmer_response
                        else "-"
                    ),

                "Acción final":
                    self._text(
                        latest.get(
                            "final_action"
                        )
                    ),
            }


            decision_table = (
                document.add_table(
                    rows=1,
                    cols=2,
                )
            )


            decision_table.rows[
                0
            ].cells[
                0
            ].text = "Campo"

            decision_table.rows[
                0
            ].cells[
                1
            ].text = "Valor"


            for key, value in (
                decision_values.items()
            ):

                cells = (
                    decision_table
                    .add_row()
                    .cells
                )

                cells[
                    0
                ].text = str(
                    key
                )

                cells[
                    1
                ].text = (
                    self._text(
                        value
                    )
                )


            self._configure_table(
                decision_table,
                8,
            )

        else:

            document.add_paragraph(
                (
                    "No existen decisiones "
                    "registradas."
                )
            )


        # =================================================
        # 4. EVIDENCE
        # =================================================

        document.add_heading(
            "4. Evidencia",
            level=1,
        )


        evidence_table = (
            document.add_table(
                rows=1,
                cols=7,
            )
        )


        headers = [
            "Variable",
            "Valor",
            "Unidad",
            "Estado",
            "Calidad",
            "Origen",
            "Fecha",
        ]


        for index, header in (
            enumerate(
                headers
            )
        ):

            evidence_table.rows[
                0
            ].cells[
                index
            ].text = header


        for item in (
            report.get(
                "evidence",
                [],
            )
        ):

            cells = (
                evidence_table
                .add_row()
                .cells
            )


            values = [

                item.get(
                    "variable"
                ),

                self._text(
                    item.get(
                        "value"
                    )
                ),

                item.get(
                    "unit"
                ),

                item.get(
                    "status"
                ),

                item.get(
                    "quality_score"
                ),

                self._origin(
                    item
                ),

                item.get(
                    "observed_at"
                ),
            ]


            for index, value in (
                enumerate(
                    values
                )
            ):

                cells[
                    index
                ].text = (
                    self._text(
                        value
                    )
                )


        self._configure_table(
            evidence_table,
            7,
        )


        # =================================================
        # 5. FARMER KNOWLEDGE
        # =================================================

        document.add_heading(
            (
                "5. Conocimiento "
                "del agricultor"
            ),
            level=1,
        )


        knowledge_table = (
            document.add_table(
                rows=1,
                cols=6,
            )
        )


        headers = [
            "Tipo",
            "Fuente",
            "Rol",
            "Contenido",
            "Validación",
            "Confianza",
        ]


        for index, header in (
            enumerate(
                headers
            )
        ):

            knowledge_table.rows[
                0
            ].cells[
                index
            ].text = header


        for item in (
            report.get(
                "farmer_knowledge",
                [],
            )
        ):

            cells = (
                knowledge_table
                .add_row()
                .cells
            )


            values = [

                item.get(
                    "knowledge_type"
                ),

                item.get(
                    "source_id"
                ),

                item.get(
                    "source_role"
                ),

                self._text(
                    item.get(
                        "payload"
                    )
                ),

                item.get(
                    "validation_status"
                ),

                item.get(
                    "reported_confidence"
                ),
            ]


            for index, value in (
                enumerate(
                    values
                )
            ):

                cells[
                    index
                ].text = (
                    self._text(
                        value
                    )
                )


        self._configure_table(
            knowledge_table,
            7,
        )


        # =================================================
        # 6. GOVERNANCE
        # =================================================

        document.add_heading(
            "6. Gobernanza",
            level=1,
        )


        governance = (
            report.get(
                "governance",
                {},
            )
        )


        governance_table = (
            document.add_table(
                rows=1,
                cols=2,
            )
        )


        governance_table.rows[
            0
        ].cells[
            0
        ].text = "Indicador"

        governance_table.rows[
            0
        ].cells[
            1
        ].text = "Valor"


        governance_values = {

            "Consentimientos totales":
                governance.get(
                    "consent_count"
                ),

            "Consentimientos activos":
                governance.get(
                    "active_consent_count"
                ),

            "Eventos de auditoría":
                governance.get(
                    "audit_event_count"
                ),
        }


        for key, value in (
            governance_values.items()
        ):

            cells = (
                governance_table
                .add_row()
                .cells
            )

            cells[
                0
            ].text = key

            cells[
                1
            ].text = (
                self._text(
                    value
                )
            )


        self._configure_table(
            governance_table,
            9,
        )


        # =================================================
        # AUDIT
        # =================================================

        events = (
            governance.get(
                "latest_audit_events",
                [],
            )
        )


        if events:

            document.add_heading(
                (
                    "Últimos eventos "
                    "de auditoría"
                ),
                level=2,
            )


            audit_table = (
                document.add_table(
                    rows=1,
                    cols=5,
                )
            )


            headers = [
                "Fecha",
                "Actor",
                "Acción",
                "Permitido",
                "Consentimiento",
            ]


            for index, header in (
                enumerate(
                    headers
                )
            ):

                audit_table.rows[
                    0
                ].cells[
                    index
                ].text = header


            for item in events:

                cells = (
                    audit_table
                    .add_row()
                    .cells
                )


                values = [

                    item.get(
                        "created_at"
                    ),

                    item.get(
                        "actor_id"
                    ),

                    item.get(
                        "action"
                    ),

                    item.get(
                        "allowed"
                    ),

                    item.get(
                        "consent_id"
                    ),
                ]


                for index, value in (
                    enumerate(
                        values
                    )
                ):

                    cells[
                        index
                    ].text = (
                        self._text(
                            value
                        )
                    )


            self._configure_table(
                audit_table,
                7,
            )

        # =================================================
        # 7. SENSORS
        # =================================================

        document.add_heading(
            "7. Sensores registrados",
            level=1,
        )


        sensors = (
            report.get(
                "sensors",
                [],
            )
        )


        if sensors:

            table = document.add_table(
                rows=1,
                cols=7,
            )

            headers = [
                "ID",
                "Nombre",
                "Tipo",
                "Variables",
                "Activo",
                "Fabricante / modelo",
                "Ubicación",
            ]


            for index, header in enumerate(
                headers
            ):
                table.rows[
                    0
                ].cells[
                    index
                ].text = header


            for sensor in sensors:

                cells = (
                    table.add_row()
                    .cells
                )

                values = [
                    sensor.get(
                        "id"
                    ),

                    sensor.get(
                        "name"
                    ),

                    sensor.get(
                        "sensor_type"
                    ),

                    ", ".join(
                        sensor.get(
                            "variables",
                            [],
                        )
                    ),

                    (
                        "Sí"
                        if sensor.get(
                            "is_active"
                        )
                        else "No"
                    ),

                    (
                        f"{sensor.get('manufacturer') or '-'} "
                        f"/ "
                        f"{sensor.get('model') or '-'}"
                    ),

                    (
                        f"{sensor.get('latitude') or '-'}, "
                        f"{sensor.get('longitude') or '-'}"
                    ),
                ]


                for index, value in enumerate(
                    values
                ):
                    cells[
                        index
                    ].text = self._text(
                        value
                    )


            self._configure_table(
                table,
                7,
            )

        else:

            document.add_paragraph(
                "No existen sensores registrados."
            )


        # =================================================
        # 8. CONSTRAINTS
        # =================================================

        document.add_heading(
            (
                "8. Restricciones aportadas "
                "por el agricultor"
            ),
            level=1,
        )


        constraints = [
            item
            for item
            in report.get(
                "farmer_knowledge",
                [],
            )
            if item.get(
                "knowledge_type"
            ) == "CONSTRAINT"
        ]


        if constraints:

            table = document.add_table(
                rows=1,
                cols=5,
            )

            headers = [
                "Tipo",
                "Afecta",
                "Contenido",
                "Justificación",
                "Validación",
            ]


            for index, header in enumerate(
                headers
            ):
                table.rows[
                    0
                ].cells[
                    index
                ].text = header


            for item in constraints:

                cells = (
                    table.add_row()
                    .cells
                )

                values = [
                    item.get(
                        "constraint_type"
                    ),

                    ", ".join(
                        item.get(
                            "affects",
                            [],
                        )
                    ),

                    self._text(
                        item.get(
                            "payload"
                        )
                    ),

                    item.get(
                        "rationale"
                    ),

                    item.get(
                        "validation_status"
                    ),
                ]


                for index, value in enumerate(
                    values
                ):
                    cells[
                        index
                    ].text = self._text(
                        value
                    )


            self._configure_table(
                table,
                7,
            )

        else:

            document.add_paragraph(
                (
                    "No existen restricciones "
                    "participativas registradas."
                )
            )


        # =================================================
        # 9. DECISION TRACE
        # =================================================

        document.add_heading(
            (
                "9. Trazabilidad del motor "
                "de inferencia y decisión"
            ),
            level=1,
        )


        document.add_paragraph(
            (
                "Esta sección pertenece al motor de "
                "inferencia y decisión. El Random "
                "Forest complementario no participa "
                "en el ranking ni modifica la "
                "recomendación."
            )
        )


        traces = report.get(
            "decision_traces",
            [],
        )


        if traces:

            latest_trace = traces[0]

            table = document.add_table(
                rows=1,
                cols=2,
            )

            table.rows[
                0
            ].cells[
                0
            ].text = "Componente"

            table.rows[
                0
            ].cells[
                1
            ].text = "Contenido persistido"


            trace_values = {

                "Trace ID":
                    latest_trace.get(
                        "trace_id"
                    ),

                "Evidencias":
                    latest_trace.get(
                        "evidence_ids",
                        [],
                    ),

                "Transformaciones":
                    latest_trace.get(
                        "transformation_ids",
                        [],
                    ),

                "Conocimiento":
                    latest_trace.get(
                        "knowledge_ids",
                        [],
                    ),

                "Conflictos":
                    latest_trace.get(
                        "conflict_ids",
                        [],
                    ),

                "Escenarios factibles":
                    latest_trace.get(
                        "feasible_scenarios",
                        [],
                    ),

                "Escenarios rechazados":
                    latest_trace.get(
                        "rejected_scenarios",
                        [],
                    ),

                "Ranking multicriterio":
                    latest_trace.get(
                        "ranking",
                        [],
                    ),

                "Recomendación provisional":
                    latest_trace.get(
                        "provisional_recommendation"
                    ),

                "Respuesta HITL":
                    latest_trace.get(
                        "farmer_response"
                    ),

                "Acción final":
                    latest_trace.get(
                        "final_action"
                    ),

                "Outcome":
                    latest_trace.get(
                        "outcome"
                    ),
            }


            for key, value in (
                trace_values.items()
            ):

                cells = (
                    table.add_row()
                    .cells
                )

                cells[
                    0
                ].text = key

                cells[
                    1
                ].text = self._text(
                    value
                )


            self._configure_table(
                table,
                8,
            )

        else:

            document.add_paragraph(
                (
                    "No existen DecisionTrace "
                    "registradas."
                )
            )


        # =================================================
        # 10. COMPLEMENTARY ML
        # =================================================

        document.add_heading(
            (
                "10. Random Forest "
                "complementario"
            ),
            level=1,
        )


        ml = report.get(
            "ml_complementary",
            {},
        )

        forecast = (
            ml.get(
                "forecast"
            )
            or {}
        )

        uncertainty = (
            forecast.get(
                "uncertainty"
            )
            or {}
        )


        table = document.add_table(
            rows=1,
            cols=2,
        )

        table.rows[
            0
        ].cells[
            0
        ].text = "Campo"

        table.rows[
            0
        ].cells[
            1
        ].text = "Valor"


        ml_values = {

            "Rol":
                (
                    "Predicción complementaria "
                    "sin autoridad decisional"
                ),

            "Modelo":
                ml.get(
                    "model"
                ),

            "Estado del modelo":
                ml.get(
                    "model_status"
                ),

            "Pronóstico disponible":
                ml.get(
                    "forecast_available"
                ),

            "Horizonte (h)":
                ml.get(
                    "prediction_horizon_hours"
                ),

            "Humedad actual":
                forecast.get(
                    "current_soil_moisture"
                ),

            "Humedad prevista +24 h":
                forecast.get(
                    "predicted_soil_moisture_24h"
                ),

            "P10":
                uncertainty.get(
                    "p10"
                ),

            "P90":
                uncertainty.get(
                    "p90"
                ),

            "Tendencia":
                forecast.get(
                    "trend"
                ),

            "Modifica TrustedState":
                False,

            "Modifica recomendación":
                False,

            "Mensaje":
                ml.get(
                    "message"
                ),
        }


        for key, value in (
            ml_values.items()
        ):

            cells = (
                table.add_row()
                .cells
            )

            cells[
                0
            ].text = key

            cells[
                1
            ].text = self._text(
                value
            )


        self._configure_table(
            table,
            8,
        )

        buffer = BytesIO()

        document.save(
            buffer
        )


        return (
            buffer.getvalue()
        )