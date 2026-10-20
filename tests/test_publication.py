from io import BytesIO
from zipfile import ZipFile
import json
import pytest
from PIL import Image
from services.publication_service import PublicationService


def sample():
    return {'parcel_id': 'P1', 'generated_at': '2026-10-20', 'evidence': [
        {'variable': 'humidity', 'unit': '%', 'value': 20, 'observed_at': '2026-10-19T00:00:00Z', 'status': 'ACTIVE'},
        {'variable': 'humidity', 'unit': '%', 'value': 30, 'observed_at': '2026-10-20T00:00:00Z', 'status': 'ACTIVE'},
        {'variable': 'humidity', 'unit': 'fraction', 'value': 0.3, 'observed_at': '2026-10-20T00:00:00Z'},
        {'variable': 'humidity', 'unit': '%', 'value': None, 'observed_at': 'bad'},
    ]}


def test_groups_units_and_explains_actual_values():
    data = PublicationService().describe(sample())
    assert len(data['figures']) == 2
    figure = data['figures'][0]
    assert figure['count'] == 2
    assert '20' in figure['interpretation'] and '30' in figure['interpretation']
    assert figure['explanation'] and figure['limitations']


def test_package_preserves_vector_and_print_dimensions():
    content = PublicationService().generate(sample(), width_mm=85)
    with ZipFile(BytesIO(content)) as archive:
        assert {'manifest.json', 'evidence.csv', 'tables.xlsx', 'captions.txt'} <= set(archive.namelist())
        metadata = json.loads(archive.read('manifest.json'))
        assert metadata['width_mm'] == 85
        png = Image.open(BytesIO(archive.read('figure-1.png')))
        assert abs(png.width - 85 / 25.4 * 600) < 2
        assert abs(png.info['dpi'][0] - 600) < 1
        assert b'<svg' in archive.read('figure-1.svg')
        assert archive.read('figure-1.pdf').startswith(b'%PDF')


def test_empty_and_invalid_width():
    assert PublicationService().describe({'evidence': []})['figures'] == []
    with pytest.raises(ValueError):
        PublicationService().generate(sample(), width_mm=0)


def test_admin_routes_reject_non_admin_before_reading_parcel():
    from fastapi import HTTPException
    from types import SimpleNamespace
    from api.routers import reports
    from api.auth_dependencies import require_admin
    for path in ('/parcels/{parcel_id}/publication', '/parcels/{parcel_id}/publication.zip'):
        route = next(route for route in reports.router.routes if route.path.endswith(path))
        assert any(dependency.call is require_admin for dependency in route.dependant.dependencies)
    with pytest.raises(HTTPException) as exc:
        require_admin(SimpleNamespace(system_role='FARMER'))
    assert exc.value.status_code == 403


def test_admin_route_checks_consent(monkeypatch):
    from fastapi import HTTPException
    from types import SimpleNamespace
    from api.routers import reports
    user = SimpleNamespace(system_role='ADMIN')
    monkeypatch.setattr(reports, '_ensure_parcel', lambda **kw: SimpleNamespace(name='P1'))
    permissions = []
    monkeypatch.setattr(reports, 'check_parcel_permission', lambda **kw: permissions.append(kw['permission']))
    monkeypatch.setattr(reports, '_build_report', lambda **kw: sample())
    def deny(**kw):
        raise HTTPException(status_code=403, detail='Consent required')
    monkeypatch.setattr(reports, '_check_export_consent', deny)
    assert len(reports.publication_preview('P1', current_user=user, db=None)['figures']) == 2
    with pytest.raises(HTTPException) as exc:
        reports.publication_download('P1', width_mm=85, current_user=user, db=None)
    assert exc.value.status_code == 403
    assert permissions[-1] == reports.ParcelPermission.EXPORT_DATA


def test_admin_report_tables_have_explanations():
    from test_report_generation import build_report
    from services.report_docx_service import ReportDOCXService
    from services.report_excel_service import ReportExcelService
    from docx import Document
    from openpyxl import load_workbook
    report = build_report()
    report['publication_admin'] = True
    document = Document(BytesIO(ReportDOCXService().generate(report)))
    for table in document.tables:
        assert 'Interpretación' in ''.join(table._tbl.getnext().itertext())
    workbook = load_workbook(BytesIO(ReportExcelService().generate(report)))
    for sheet in workbook:
        assert any('Interpretación' in str(cell.value) for row in sheet for cell in row)


def test_pdf_explanations_follow_each_admin_table(monkeypatch):
    from test_report_generation import build_report
    from services.report_pdf_service import ReportPDFService
    from reportlab.platypus import SimpleDocTemplate, Table, Paragraph
    story = []
    monkeypatch.setattr(SimpleDocTemplate, 'build', lambda self, content: story.extend(content))
    report = build_report()
    report['publication_admin'] = True
    ReportPDFService().generate(report)
    for i, element in enumerate(story):
        if isinstance(element, Table):
            assert isinstance(story[i + 1], Paragraph)
            assert 'Interpretación' in story[i + 1].text


def test_excel_preserves_numbers_and_escapes_formulas():
    from openpyxl import load_workbook
    report = sample()
    report['evidence'][0]['provenance'] = '=HYPERLINK("bad")'
    with ZipFile(BytesIO(PublicationService().generate(report))) as archive:
        sheet = load_workbook(BytesIO(archive.read('tables.xlsx'))).active
        assert sheet['B2'].value == 20
        assert sheet['H2'].value.startswith("'=")
