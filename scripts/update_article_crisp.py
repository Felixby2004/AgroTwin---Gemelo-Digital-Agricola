"""Local OOXML edit preserving the supplied journal manuscript."""
from copy import deepcopy
from pathlib import Path
import json
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import csv

ROOT=Path(__file__).resolve().parents[1]
SOURCE=Path(r'C:\Users\FELIX\Desktop\CICLO VIII\ING. SOFTWARE 2\AgroTwin_ESPANOL.docx')
ASSETS=ROOT/'artifacts/crisp_publicacion'
OUT=ROOT/'artifacts/AgroTwin_ESPANOL_CRISP_DM_actualizado.docx'
d=Document(SOURCE)
original=d.paragraphs[:]
body_ref=original[79]
caption_ref=original[78]


def style_paragraph(p, reference):
    p.style=reference.style
    if reference._p.pPr is not None:
        if p._p.pPr is not None: p._p.remove(p._p.pPr)
        p._p.insert(0,deepcopy(reference._p.pPr))


def set_text(p,text):
    rpr=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear()
    r=p.add_run(text)
    if rpr is not None: r._r.insert(0,rpr)


def after(anchor,text,reference=None):
    element=OxmlElement('w:p')
    anchor._p.addnext(element)
    paragraph=Paragraph(element,anchor._parent)
    reference=reference or body_ref
    style_paragraph(paragraph,reference)
    r=paragraph.add_run(text)
    if reference.runs and reference.runs[0]._r.rPr is not None:
        r._r.insert(0,deepcopy(reference.runs[0]._r.rPr))
    return paragraph


def figure_after(anchor,key,caption):
    image=after(anchor,'',body_ref)
    image.alignment=WD_ALIGN_PARAGRAPH.CENTER
    image.paragraph_format.keep_with_next=True
    image.paragraph_format.keep_together=True
    shape=image.add_run().add_picture(str(ASSETS/(key+'.png')),width=Mm(85))
    shape._inline.docPr.set('descr',caption)
    cap=after(image,caption,caption_ref)
    cap.paragraph_format.keep_together=True
    cap.paragraph_format.keep_with_next=False
    cap.paragraph_format.space_after=Pt(6)
    return cap


def replace_figure(p,key,caption,p_caption):
    p.clear()
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.keep_together=True
    shape=p.add_run().add_picture(str(ASSETS/(key+'.png')),width=Mm(85))
    shape._inline.docPr.set('descr',caption)
    set_text(p_caption,caption)
    p_caption.paragraph_format.keep_together=True


# Preserve numbering in reading order after the three new CRISP-DM figures.
for p in original:
    if 'Fig. 6' in p.text: set_text(p,p.text.replace('Fig. 6','Fig. 9'))
    if p.text.startswith('TABLA VI.'): set_text(p,p.text.replace('TABLA VI.','TABLA VII.',1))

replace_figure(original[77],'model_selection_figure',
    'Fig. 4. Comparación de MAE y RMSE de los candidatos en 2024, generada por el laboratorio CRISP-DM de AgroTwin. Menor error indica mejor ajuste predictivo. Todos los candidatos comparten partición temporal y 25 predictores; la selección utiliza MAE sin redondear. La figura no demuestra eficacia del riego.',original[78])
replace_figure(original[81],'bootstrap',
    'Fig. 5. Diferencia Extra Trees menos persistencia en 2025. Los intervalos del 95 % se obtienen mediante 2000 remuestreos por bloques diarios. El IC del MAE cruza cero; el del RMSE permanece negativo y favorece al modelo seleccionado. La dependencia temporal se trata mediante bloques, no como errores horarios independientes.',original[82])

p=after(original[69],
    'Las figuras predictivas y de explicabilidad se regeneraron con las mismas funciones de ml.reporting utilizadas por el panel app_streamlit.crisp_dm_panel. Se conservaron los artefactos de evaluación y se ajustó únicamente la presentación editorial a 85 mm por columna, con PNG a 600 dpi y copias SVG y PDF vectoriales. La importancia por permutación emplea 3000 observaciones reservadas de 2025 y 10 repeticiones; SHAP utiliza 500 observaciones del mismo periodo. Ambos muestreos fijan la semilla 42.')
after(p,
    'De forma separada se preparó un conjunto demostrativo para la aplicación: cuatro parcelas ficticias, 12 sensores simulados, 6740 lecturas horarias, 24 aportes de conocimiento y cuatro trazas calculadas con los escenarios de referencia del motor. Los registros se identifican como SYNTHETIC. Este conjunto permite inspeccionar los flujos del software; no participa en el entrenamiento ni en las métricas predictivas NOAA y no representa validación de campo.')

p=after(original[84],
    'La Fig. 6 complementa las métricas con la trayectoria temporal de la humedad observada y pronosticada. La agregación diaria permite identificar cambios estacionales y episodios con mayor separación entre curvas; no sustituye el cálculo de MAE y RMSE sobre los registros horarios reservados.')
figure_after(p,'final_timeseries',
    'Fig. 6. Humedad del suelo observada, predicha por Extra Trees y estimada por persistencia en 2025. Se muestran medias diarias de las predicciones horarias guardadas por CRISP-DM. La cercanía de las curvas describe seguimiento temporal; la agregación puede ocultar errores horarios y no permite inferir causalidad ni calidad del riego.')

p=after(original[86],
    'Las Figs. 7 y 8 muestran las salidas efectivas de explicabilidad del laboratorio. La permutación sitúa primero la humedad actual del suelo (incremento medio de MAE de 0.000982 m³/m³) y la humedad relativa (0.000830 m³/m³). SHAP asigna las mayores contribuciones absolutas medias a la humedad relativa (0.000828 m³/m³) y a la humedad actual (0.000790 m³/m³). Los órdenes diferentes son compatibles con predictores correlacionados y con métodos que cuantifican aspectos distintos de la influencia predictiva.')
p=figure_after(p,'permutation_importance',
    'Fig. 7. Quince variables principales por importancia de permutación en el conjunto reservado de 2025. Las barras muestran el aumento medio de MAE al desordenar cada predictor y las marcas negras su desviación estándar en 10 repeticiones sobre 3000 observaciones. Mayor aumento implica mayor dependencia predictiva; no demuestra una relación causal.')
p=figure_after(p,'shap_global',
    'Fig. 8. Importancia global SHAP del predictor a +24 h, calculada sobre 500 observaciones reservadas de 2025. Cada barra representa la media del valor absoluto de la contribución, expresada en m³/m³. Destacan humedad relativa y humedad actual; el valor absoluto mide magnitud, sin identificar por sí solo dirección ni causalidad.')
p=after(p,'La Tabla VI conserva una explicación individual como ejemplo: el signo de SHAP indica el desplazamiento del cambio de humedad predicho para esa observación, mientras que la magnitud permite ordenar sus contribuciones. No se interpreta como una recomendación de riego.')
p=after(p,'TABLA VI. PRINCIPALES CONTRIBUCIONES SHAP DE UNA OBSERVACIÓN RESERVADA',original[83])
p.paragraph_format.keep_with_next=True
rows=list(csv.DictReader((ASSETS/'shap_local.csv').open(encoding='utf-8-sig')))[:5]
t=d.add_table(rows=1,cols=3)
t.style=d.tables[4].style
t.autofit=False
for cell,label in zip(t.rows[0].cells,['Variable','Entrada','SHAP ×10⁻⁴']): cell.text=label
names={'soil_moisture':'Humedad actual','relative_humidity':'Hum. relativa','humidity_6h_mean':'Hum. media 6 h','doy_cos':'Día del año cos.','doy_sin':'Día del año seno','hour_cos':'Hora coseno','soil_moisture_lag_24h':'Humedad t−24 h','rainfall_24h':'Lluvia 24 h'}
for row in rows:
    cells=t.add_row().cells
    cells[0].text=names.get(row['feature'],row['feature'])
    cells[1].text=f"{float(row['input_value']):.5g}"
    cells[2].text=f"{float(row['shap_value'])*1e4:+.4f}"
for row in t.rows:
    for i,cell in enumerate(row.cells):
        cell.width=Mm([38,20,27][i])
        for para in cell.paragraphs:
            para.paragraph_format.space_after=Pt(2)
            para.paragraph_format.space_before=Pt(2)
            for run in para.runs:
                run.font.name='Times New Roman'
                run.font.size=Pt(8)
# Journal table rules with a repeated bold header.
borders=OxmlElement('w:tblBorders')
for edge in ('top','bottom'):
    rule=OxmlElement('w:'+edge);rule.set(qn('w:val'),'single');rule.set(qn('w:sz'),'4');rule.set(qn('w:color'),'000000');borders.append(rule)
t._tbl.tblPr.append(borders)
for cell in t.rows[0].cells:
    for para in cell.paragraphs:
        for run in para.runs: run.bold=True
# Prevent splitting data rows and repeat header if needed.
for row in t.rows:
    trpr=row._tr.get_or_add_trPr()
    trpr.append(OxmlElement('w:cantSplit'))
t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
p._p.addnext(t._tbl)
notes=OxmlElement('w:p'); t._tbl.addnext(notes)
np=Paragraph(notes,p._parent);style_paragraph(np,caption_ref)
np.add_run('Nota. Fuente: tabla SHAP local generada por ml.reporting para la observación más reciente de la muestra de 500 registros de 2025 (semilla 42). Un signo negativo reduce el cambio pronosticado y uno positivo lo incrementa. Las entradas conservan sus unidades originales; la codificación del día del año es adimensional. Las contribuciones no constituyen efectos causales.')
for run in np.runs:run.font.size=Pt(8)

# Clarify the exact cross-validation comparison from the saved table.
set_text(original[79],original[79].text.replace('El MAE medio en ventanas expansivas es aproximadamente 0.0037 tanto para persistencia como para Extra Trees', 'El MAE medio en ventanas expansivas es 0.003666 para persistencia y 0.003732 para Extra Trees'))
set_text(original[117],original[117].text+' Las figuras CRISP-DM actualizadas se acompañan de tablas CSV, archivos SVG y PDF vectoriales, PNG a 600 dpi y un manifiesto con las huellas SHA-256 de los artefactos utilizados. Los datos demostrativos de Neon se distribuyen por separado en un SQL idempotente identificado como sintético.')
set_text(original[123],original[123].text+' La actualización editorial y los scripts de exportación contaron con asistencia de ChatGPT; las figuras CRISP-DM se renderizaron mediante Matplotlib a partir de las salidas del sistema, sin síntesis de imágenes por un modelo generativo.')

# Keep author portraits with their own biographies instead of floating into another column.
for author_index in (166,169):
    title=original[author_index]
    for drawing in list(title._p.xpath('.//w:drawing')):
        portrait=after(title,'',body_ref)
        portrait.alignment=WD_ALIGN_PARAGRAPH.LEFT
        portrait.paragraph_format.keep_with_next=True
        portrait.paragraph_format.keep_together=True
        anchor=drawing.find(qn('wp:anchor'))
        if anchor is not None:
            anchor.tag=qn('wp:inline')
            for child in list(anchor):
                if child.tag in {qn('wp:'+tag) for tag in ('simplePos','positionH','positionV','wrapNone','wrapSquare','wrapTopAndBottom','wrapThrough','wrapTight')}:
                    anchor.remove(child)
            for attribute in list(anchor.attrib):
                if attribute not in ('distT','distB','distL','distR'): del anchor.attrib[attribute]
        portrait.add_run()._r.append(drawing)
    title.paragraph_format.keep_with_next=True
# Keep multiline equations and table titles within one column/page.
for equation in original[46:53]:
    equation.paragraph_format.keep_together=True
for paragraph in d.paragraphs:
    if paragraph.text.startswith('TABLA '):
        paragraph.paragraph_format.keep_together=True
        paragraph.paragraph_format.keep_with_next=True
# Keep the source journal layout and all equations, references, author data, and other figures.
d.save(OUT)
changes={'source':str(SOURCE),'output':str(OUT),'replaced_figures':[4,5],'added_figures':[6,7,8],'renumbered_stability_figure':9,'added_table':'VI SHAP local','synthetic_table_renumbered':'VII','original_figures_retained':[1,2,3,9],'neon_synthetic_data_used_for_crisp_metrics':False}
(ROOT/'artifacts/crisp_publicacion/article_changes.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),encoding='utf-8')
print(str(OUT))
